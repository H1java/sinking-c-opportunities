# -*- coding: utf-8 -*-
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from check_data import DIRECTIONS, META

OUT = "artifacts"
os.makedirs(OUT, exist_ok=True)

DARK = PatternFill("solid", fgColor="1F4E79")
D1 = PatternFill("solid", fgColor="DDEBF7")
D2 = PatternFill("solid", fgColor="FCE4D6")
D3 = PatternFill("solid", fgColor="E2EFDA")
RED = PatternFill("solid", fgColor="FFC7CE")
WHITE_F = Font(color="FFFFFF", bold=True, size=11)
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
FILLS = {"方向①": D1, "方向②": D2, "方向③": D3}

wb = Workbook()

def header(ws, headers, widths):
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.fill = DARK; c.font = WHITE_F; c.alignment = CENTER; c.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"

# 1. 核查总览
ws = wb.active; ws.title = "核查总览"
header(ws, ["方向", "赛道", "目标人群", "赛道阶段", "进入空间结论", "结论依据"], [8, 26, 22, 34, 24, 80])
for r, d in enumerate(DIRECTIONS, 2):
    vals = [d["id"], d["name"], d["audience"], d["stage"], d["entry"], d["entry_basis"]]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.border = THIN
        c.alignment = WRAP if i in (2, 3, 4, 5, 6) else CENTER
        if i == 1: c.fill = FILLS[d["id"]]
        if i == 5: c.fill = RED
    ws.row_dimensions[r].height = 150
r = len(DIRECTIONS) + 3
ws.cell(r, 1, "通道说明：" + META["channel_note"]).font = Font(italic=True, color="808080")
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)

# 2. 同行明细
ws = wb.create_sheet("同行明细")
header(ws, ["方向", "账号/店铺", "平台", "粉丝/销量", "更新频率", "内容形式", "发布时间", "可见销量/收入", "URL", "备注"],
       [8, 26, 18, 28, 18, 26, 14, 30, 40, 34])
r = 2
for d in DIRECTIONS:
    start = r
    for row in d["competitors"]:
        vals = ["", row[0], row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = THIN; c.alignment = WRAP if i in (2, 4, 6, 8, 10) else CENTER
        c = ws.cell(r, 9, row[7]); c.hyperlink = row[7]; c.font = Font(color="0563C1", underline="single")
        r += 1
    ws.cell(start, 1, d["id"]).fill = FILLS[d["id"]]
    ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    ws.cell(start, 1).alignment = CENTER
    note_r = r
    ws.cell(note_r, 2, "说明：" + d["competitor_note"]).font = Font(italic=True, color="808080")
    ws.merge_cells(start_row=note_r, start_column=2, end_row=note_r, end_column=10)
    ws.cell(note_r, 2).alignment = WRAP
    r += 1

# 3. 评论区反馈
ws = wb.create_sheet("评论区反馈")
header(ws, ["方向", "类型", "原话/观察", "URL"], [8, 18, 100, 45])
r = 2
for d in DIRECTIONS:
    start = r
    for typ, quote, url in d["comments"]:
        ws.cell(r, 2, typ).alignment = WRAP; ws.cell(r, 2).border = THIN
        ws.cell(r, 3, quote).alignment = WRAP; ws.cell(r, 3).border = THIN
        c = ws.cell(r, 4, url); c.border = THIN; c.hyperlink = url; c.font = Font(color="0563C1", underline="single")
        r += 1
    ws.cell(start, 1, d["id"]).fill = FILLS[d["id"]]
    ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    ws.cell(start, 1).alignment = CENTER

# 4. 变现实况
ws = wb.create_sheet("变现实况")
header(ws, ["方向", "收费方式", "价格", "证据", "证据等级", "URL"], [8, 30, 26, 50, 20, 45])
r = 2
for d in DIRECTIONS:
    start = r
    for method, price, ev, level, url in d["monetization"]:
        vals = ["", method, price, ev, level, url]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = THIN; c.alignment = WRAP if i in (2, 3, 4, 5) else CENTER
        c = ws.cell(r, 6, url); c.hyperlink = url; c.font = Font(color="0563C1", underline="single")
        r += 1
    ws.cell(start, 1, d["id"]).fill = FILLS[d["id"]]
    ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    ws.cell(start, 1).alignment = CENTER

# 5. 专项核实
ws = wb.create_sheet("专项核实")
header(ws, ["方向", "核实项", "结论", "证据等级", "URL"], [8, 26, 95, 22, 45])
r = 2
for d in DIRECTIONS:
    start = r
    for item, concl, level, url in d["special"]:
        vals = ["", item, concl, level, url]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = THIN; c.alignment = WRAP if i in (2, 3, 4) else CENTER
        c = ws.cell(r, 5, url); c.hyperlink = url; c.font = Font(color="0563C1", underline="single")
        r += 1
    ws.cell(start, 1, d["id"]).fill = FILLS[d["id"]]
    ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    ws.cell(start, 1).alignment = CENTER

# 6. 阶段与差异化
ws = wb.create_sheet("阶段与差异化")
header(ws, ["方向", "赛道阶段", "阶段依据", "差异化建议"], [8, 34, 70, 70])
r = 2
for d in DIRECTIONS:
    ws.cell(r, 1, d["id"]).fill = FILLS[d["id"]]; ws.cell(r, 1).alignment = CENTER; ws.cell(r, 1).border = THIN
    ws.cell(r, 2, d["stage"]).alignment = WRAP; ws.cell(r, 2).border = THIN
    ws.cell(r, 3, d["stage_basis"]).alignment = WRAP; ws.cell(r, 3).border = THIN
    ws.cell(r, 4, "\n".join("• " + x for x in d["differentiation"])).alignment = WRAP; ws.cell(r, 4).border = THIN
    ws.row_dimensions[r].height = 200
    r += 1

# 7. 陷阱与风险
ws = wb.create_sheet("陷阱与风险")
header(ws, ["方向", "陷阱/风险"], [8, 130])
r = 2
for d in DIRECTIONS:
    start = r
    for t in d["traps"]:
        ws.cell(r, 2, "• " + t).alignment = WRAP; ws.cell(r, 2).border = THIN
        r += 1
    ws.cell(start, 1, d["id"]).fill = FILLS[d["id"]]
    ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    ws.cell(start, 1).alignment = CENTER

xlsx_path = os.path.join(OUT, "下沉C端前3方向同行核查表.xlsx")
wb.save(xlsx_path)
print("saved:", xlsx_path)

# ---- Markdown ----
md = ["# 下沉C端前3方向同行核查表（2026-10）\n",
      "> " + META["channel_note"], "",
      "> 方法：" + META["method"], ""]
md.append("## 一、核查总览\n")
md.append("| 方向 | 赛道 | 阶段 | 进入空间 |")
md.append("|---|---|---|---|")
for d in DIRECTIONS:
    md.append(f"| {d['id']} | {d['name']} | {d['stage']} | **{d['entry']}** |")
md.append("")
for d in DIRECTIONS:
    md.append(f"## {d['id']} {d['name']}\n")
    md.append(f"- **目标人群**：{d['audience']}")
    md.append(f"- **赛道阶段**：{d['stage']}")
    md.append(f"- **阶段依据**：{d['stage_basis']}")
    md.append(f"- **进入空间结论**：{d['entry']}")
    md.append(f"- **结论依据**：{d['entry_basis']}")
    md.append("\n### 同行明细")
    md.append("| 账号/店铺 | 平台 | 粉丝/销量 | 更新频率 | 内容形式 | 可见销量/收入 | URL |")
    md.append("|---|---|---|---|---|---|---|")
    for row in d["competitors"]:
        md.append(f"| {row[0]} | {row[1]} | {row[2]} | {row[3]} | {row[4]} | {row[6]} | {row[7]} |")
    md.append(f"\n> {d['competitor_note']}")
    md.append("\n### 评论区反馈")
    for typ, quote, url in d["comments"]:
        md.append(f"- **{typ}**：{quote} — {url}")
    md.append("\n### 变现实况")
    for method, price, ev, level, url in d["monetization"]:
        md.append(f"- {method}（{price}）：{ev}〔{level}〕{url}")
    md.append("\n### 专项核实")
    for item, concl, level, url in d["special"]:
        md.append(f"- **{item}**：{concl}〔{level}〕{url}")
    md.append("\n### 差异化建议")
    for x in d["differentiation"]:
        md.append(f"- {x}")
    md.append("\n### 陷阱与风险")
    for x in d["traps"]:
        md.append(f"- {x}")
    md.append("")
md_path = os.path.join(OUT, "下沉C端前3方向同行核查表.md")
with open(md_path, "w", encoding="utf-8") as f:
    f.write("\n".join(md))
print("saved:", md_path)

# ---- 回读自检 ----
from openpyxl import load_workbook
wb2 = load_workbook(xlsx_path)
print("sheets:", wb2.sheetnames)
for s in wb2.sheetnames:
    print(s, wb2[s].max_row, wb2[s].max_column)
