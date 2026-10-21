# -*- coding: utf-8 -*-
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from cross_data import WEIGHTS, VARIABLES, CANDIDATES, REJECTED, DEATH_PATTERNS, META

OUT = "artifacts"
os.makedirs(OUT, exist_ok=True)

DARK = PatternFill("solid", fgColor="1F4E79")
GREEN = PatternFill("solid", fgColor="C6EFCE")
GRAY = PatternFill("solid", fgColor="E7E6E6")
CAT_FILL = {"模型":PatternFill("solid",fgColor="DDEBF7"),"API":PatternFill("solid",fgColor="DDEBF7"),
            "政策":PatternFill("solid",fgColor="FCE4D6"),"平台":PatternFill("solid",fgColor="E2EFDA"),
            "支付":PatternFill("solid",fgColor="FFF2CC"),"数据":PatternFill("solid",fgColor="EDEDED"),
            "Agent能力":PatternFill("solid",fgColor="DDEBF7"),"开源":PatternFill("solid",fgColor="D9E1F2"),
            "语音":PatternFill("solid",fgColor="DDEBF7"),"图像":PatternFill("solid",fgColor="DDEBF7"),
            "分发":PatternFill("solid",fgColor="E2EFDA"),"链上":PatternFill("solid",fgColor="FFF2CC")}
WHITE_F = Font(color="FFFFFF", bold=True, size=11)
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)

wb = Workbook()

def header(ws, headers, widths):
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.fill = DARK; c.font = WHITE_F; c.alignment = CENTER; c.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"

# 1. 重排榜单
ws = wb.active; ws.title = "重排榜单"
header(ws, ["排名","ID","机会","目标人群","形态","总分","一句话结论"], [6,10,30,28,34,7,60])
for r, c in enumerate(CANDIDATES, 2):
    vals = [r-1, c["id"], c["name"], c["audience"], c["form"], c["score"], c["why_now"]]
    for i, v in enumerate(vals, 1):
        cell = ws.cell(r, i, v); cell.border = THIN
        cell.alignment = WRAP if i in (3,4,5,7) else CENTER
        if i == 6: cell.fill = GREEN
    ws.row_dimensions[r].height = 70

# 2. 新变量台账
ws = wb.create_sheet("新变量台账30天")
header(ws, ["日期","类别","名称","URL"], [9,11,60,70])
for r, (date, cat, name, url) in enumerate(VARIABLES, 2):
    ws.cell(r,1,date).alignment=CENTER; ws.cell(r,1).border=THIN
    cc=ws.cell(r,2,cat); cc.alignment=CENTER; cc.border=THIN; cc.fill=CAT_FILL.get(cat,GRAY)
    ws.cell(r,3,name).alignment=WRAP; ws.cell(r,3).border=THIN
    cell=ws.cell(r,4,url); cell.border=THIN; cell.hyperlink=url; cell.font=Font(color="0563C1",underline="single")

# 3. 证据档案（通过项明细）
ws = wb.create_sheet("通过项明细")
header(ws, ["ID","机会","字段","内容"], [10,28,22,100])
r=2
for c in CANDIDATES:
    blocks=[("目标人群",c["audience"]),("形态",c["form"]),
            ("为什么是现在",c["why_now"]),("第一个客户触达",c["first_customer"]),
            ("定价",c["pricing"]),("30天验证指标",c["metric"]),
            ("止损线",c["stop_loss"]),("风险/不适合人群",c["risk"])]
    for q,e in c["new_vars"]: blocks.append(("新变量："+q,e))
    for q,e in c["evidence"]: blocks.append(("证据："+q,e))
    start=r
    for k,v in blocks:
        ws.cell(r,3,k).alignment=WRAP; ws.cell(r,3).border=THIN
        ws.cell(r,4,v).alignment=WRAP; ws.cell(r,4).border=THIN
        if k.startswith(("新变量","证据")): ws.cell(r,3).font=Font(color="C00000")
        r+=1
    ws.cell(start,1,c["id"]); ws.cell(start,2,f"{c['name']}（{c['score']}分）")
    ws.merge_cells(start_row=start,start_column=1,end_row=r-1,end_column=1)
    ws.merge_cells(start_row=start,start_column=2,end_row=r-1,end_column=2)
    ws.cell(start,1).alignment=CENTER; ws.cell(start,2).alignment=CENTER

# 4. 评分明细
ws = wb.create_sheet("评分明细")
dim_names=list(WEIGHTS.keys())
header(ws,["ID","机会","总分"]+[f"{k}({v})" for k,v in WEIGHTS.items()],[10,30,7]+[11]*7)
keys=["付费","痛点","竞品","频次","付费意愿","免费替代","薄壳"]
for r,c in enumerate(CANDIDATES,2):
    ws.cell(r,1,c["id"]).border=THIN; ws.cell(r,2,c["name"]).border=THIN; ws.cell(r,2).alignment=WRAP
    ws.cell(r,3,c["score"]).border=THIN; ws.cell(r,3).alignment=CENTER
    for i,k in enumerate(keys):
        cell=ws.cell(r,4+i,c["dim"].get(k,"-")); cell.border=THIN; cell.alignment=CENTER

# 5. 本轮淘汰台账
ws = wb.create_sheet("本轮淘汰台账")
header(ws,["构想","淘汰原因"],[40,110])
for r,(name,why) in enumerate(REJECTED,2):
    ws.cell(r,1,name).alignment=WRAP; ws.cell(r,1).border=THIN
    ws.cell(r,2,why).alignment=WRAP; ws.cell(r,2).border=THIN

# 6. 已知死亡模式
ws = wb.create_sheet("已知死亡模式")
header(ws,["#","死亡模式（反向过滤器）"],[6,130])
for r,p in enumerate(DEATH_PATTERNS,2):
    ws.cell(r,1,r-1).alignment=CENTER; ws.cell(r,1).border=THIN
    ws.cell(r,2,p).alignment=WRAP; ws.cell(r,2).border=THIN

# 7. 结论
ws = wb.create_sheet("结论")
header(ws,["要点","内容"],[22,110])
concl=[
 ("最终通过",f"{len(CANDIDATES)}个：App上架预检81、标书废标预检73、基金盲审预检70、AI购物推荐监测70、小语种成品旁白66"),
 ("核心规律", "这30天同时发生‘能力涨价/降价分化+收款通道松绑+分发入口重构’；过线的不是‘做一个AI产品’，而是‘替一个具体人群，在一个刚变的规则下，避免一个确定的损失’——止血型、证据型、可续费"),
 ("首选顺序", "1)App上架预检（规则刚改5天、付费锚点最强）2)标书废标预检（生死痛点、代理渠道可复用）3)AI购物监测（顺风、高续费但需求待证）4)基金预检（季节性）5)旁白代交付（窗口套利）"),
 ("统一风险", "5个都是薄壳，护城河=带页码/截图的证据链+垂直规则库+渠道关系；都需30天手动验证，禁止先做自动化"),
 ("通道说明", META["channel_note"]),
]
for r,(k,v) in enumerate(concl,2):
    ws.cell(r,1,k).alignment=WRAP; ws.cell(r,1).font=BOLD; ws.cell(r,1).border=THIN
    ws.cell(r,2,v).alignment=WRAP; ws.cell(r,2).border=THIN

xlsx_path=os.path.join(OUT,"最新变量交叉创新报告_2026Q4.xlsx")
wb.save(xlsx_path); print("saved:",xlsx_path)

# ---- Markdown ----
md=["# 最新变量交叉创新报告（2026 Q4）\n","> 窗口："+META["window"],"> "+META["channel_note"],"> 方法："+META["method"],""]
md.append("## 一、重排榜单\n")
md.append("| 排名 | ID | 机会 | 人群 | 总分 |")
md.append("|---|---|---|---|---|")
for i,c in enumerate(CANDIDATES,1):
    md.append(f"| {i} | {c['id']} | {c['name']} | {c['audience']} | **{c['score']}** |")
md.append("\n## 二、新变量台账（近30天）\n")
for date,cat,name,url in VARIABLES:
    md.append(f"- {date}〔{cat}〕{name} {url}")
md.append("\n## 三、通过项明细\n")
for c in CANDIDATES:
    md.append(f"### {c['id']} {c['name']}（{c['score']}分）")
    md.append(f"- 人群/形态：{c['audience']}｜{c['form']}")
    md.append(f"- 为什么是现在：{c['why_now']}")
    md.append(f"- 第一个客户：{c['first_customer']}")
    md.append(f"- 定价：{c['pricing']}")
    md.append(f"- 30天指标：{c['metric']}")
    md.append(f"- 止损：{c['stop_loss']}")
    md.append(f"- 风险：{c['risk']}")
    for q,e in c["new_vars"]: md.append(f"  - 新变量：{q} {e}")
    for q,e in c["evidence"]: md.append(f"  - 证据：{q} {e}")
    md.append("")
md.append("## 四、本轮淘汰台账\n")
for name,why in REJECTED: md.append(f"- **{name}**：{why}")
md.append("\n## 五、已知死亡模式\n")
for p in DEATH_PATTERNS: md.append(f"- {p}")
md.append("")
md_path=os.path.join(OUT,"最新变量交叉创新报告_2026Q4.md")
with open(md_path,"w",encoding="utf-8") as f: f.write("\n".join(md))
print("saved:",md_path)

# ---- 回读自检 ----
from openpyxl import load_workbook
wb2=load_workbook(xlsx_path)
print("sheets:",wb2.sheetnames)
for s in wb2.sheetnames: print(s,wb2[s].max_row,wb2[s].max_column)
